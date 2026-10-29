"""R&D 패스트트랙 통합 실행기 (pipeline_runner.py 수정본)

이번 수정
 1. 순서: 의도 분류(규칙 또는 LLM 라우터) → 통관 질문으로 확정된 뒤에만 수량·목적·단가 규칙 검사 → HSK 추론 → 보고서.
    예전에는 수량 검사가 라우터보다 앞에 있어 "안녕하세요", "연구실 소화기 점검 주기", "통관 서류가 뭐야?" 가 전부
    "수량을 파악할 수 없습니다"로 끝났고, '연구' 키워드와 수량만 있어도 LLM 없이 통관으로 분류됐다.
    이제 규칙 기반 통관 판정은 '수입 관련 표현 + 수량'이 함께 있을 때만 쓰고, 그 밖은 LLM 라우터가 판단한다.
    규칙 검사는 여전히 HSK 추론(LLM 호출) 앞에 있어서, 거절될 입력에 API 한도를 쓰지 않는다.
 2. 템플릿 모드: config.generate_text 를 통째로 바꿔치기하지 않는다. 스레드별 플래그로 '이 요청의' 사유서 LLM 호출만 건너뛰고,
    main_pipeline 의 표준 문구(품목명·수량 포함)를 쓴다. 예외가 나도 다른 요청·이후 요청에 영향이 없다.
 3. 캐시: 성공 결과만 저장·반환한다. 키는 줄바꿈(CRLF/LF) 차이와 무관한 해시 + 수동 버전(CACHE_VERSION)이다.
    Windows 와 클라우드(Linux)에서 키가 달라지던 문제를 없애고, 질문 전문이 찍히던 디버그 출력을 지웠다.
 4. LLM 장애로 키워드도 맞지 않을 때 '관련 없는 질문'이라고 잘못 안내하지 않고, 장애 문구와 error_kind 를 돌려준다.

반환값 규칙
    status     : "success" | "error" | "unsupported"
    stage      : "input" | "router" | "report"(규칙 검사 포함) | "hsk" | "done"
    error_kind : (선택) API 오류일 때 "quota" | "overloaded" | "config" | "other"
    cache_hit  : (선택) 같은 질문의 이전 성공 결과를 그대로 돌려줬을 때 True
"""
import contextlib
import hashlib
import json
import logging
import sys
import threading
import uuid
from pathlib import Path

_SRC_DIR = Path(__file__).resolve().parent
if str(_SRC_DIR) not in sys.path:
    sys.path.insert(0, str(_SRC_DIR))

import config
import report_logic as rl
from intent_router import IntentRouter, CUSTOMS, LAB_SAFETY, UNKNOWN
from hsk_inferencer import infer_hsk_code
from main_pipeline import generate_fast_track_report, CSV_PATH

logger = logging.getLogger("rnd_pipeline")
_router = None

CACHE_VERSION = "v3"  # 분석 로직·품목표 구조를 크게 바꾸면 올린다. 값이 다르면 예전 캐시는 쓰지 않는다.
CACHE_PREFIX = f"cache_{CACHE_VERSION}_"
CACHE_DIR = _SRC_DIR.parent / "cache"
try:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
except OSError:  # 읽기 전용 파일시스템이어도 앱은 캐시 없이 동작한다
    pass

# main_pipeline.generate_fast_track_report 와 같은 문구 (규칙 검사를 HSK 추론 앞에서 하므로 같은 안내를 쓴다)
QTY_MESSAGE = "수량을 확인하지 못했습니다. 예: '배터리 25개'처럼 수량을 포함해 다시 입력해 주세요."
FOREIGN_PRICE_MESSAGE = (
    "외화 단가가 입력되었습니다. 적용 환율을 알 수 없어 계산하지 않았습니다. "
    "원화로 환산한 단가(예: '개당 130만원')로 다시 입력해 주세요."
)

# 규칙 기반 통관 판정에 쓰는 '수입 관련 표현'. 수량과 함께 있을 때만 LLM 없이 통관으로 본다.
_IMPORT_TERMS = ("수입", "통관", "반입", "들여", "직구", "관세", "import")


# ── 템플릿 모드 (스레드별 플래그) ────────────────────────────────────────────
class TemplateModeSkip(Exception):
    """템플릿 모드: 사유서 LLM 호출을 건너뛴다. main_pipeline.generate_rich_reason 은 예외를 받으면 표준 문구를 쓴다."""


def _install_template_dispatcher():
    """config.generate_text 를 '평소엔 원래 함수, 템플릿 모드인 스레드에서만 예외'인 중계 함수로 한 번만 교체한다."""
    current = config.generate_text
    if getattr(current, "_rnd_dispatcher", False):  # 모듈이 다시 import 되어도 중복으로 감싸지 않는다
        return current

    local = threading.local()

    def dispatcher(prompt):
        if getattr(local, "template", False):
            raise TemplateModeSkip("템플릿 모드: 사유서 LLM 호출 생략")
        return dispatcher._rnd_real(prompt)

    dispatcher._rnd_dispatcher = True
    dispatcher._rnd_real = current
    dispatcher._rnd_local = local
    config.generate_text = dispatcher
    return dispatcher


_DISPATCHER = _install_template_dispatcher()


@contextlib.contextmanager
def _template_mode(enabled: bool):
    local = _DISPATCHER._rnd_local
    previous = getattr(local, "template", False)
    local.template = bool(enabled)
    try:
        yield
    finally:  # 예외가 나도 반드시 되돌린다
        local.template = previous


# ── 기본 도우미 ──────────────────────────────────────────────────────────────
def _get_router():
    global _router
    if _router is None:
        _router = IntentRouter()
    return _router


def _fail(stage: str, message: str, status: str = "error", **extra) -> dict:
    result = {"status": status, "stage": stage, "message": message}
    result.update(extra)
    return result


def _kind_from_message(message: str) -> str:
    """라우터 폴백 사유 문구에서 오류 종류를 되짚는다 (config.format_api_error 의 문구 기준)."""
    if "사용량 한도" in message:
        return "quota"
    if "혼잡" in message:
        return "overloaded"
    if "GEMINI_API_KEY" in message:
        return "config"
    return "other"


def _looks_like_import_request(text: str, quantity) -> bool:
    """'수입 관련 표현 + 수량'이 함께 있을 때만 규칙으로 통관 질문이라고 본다."""
    if quantity is None:
        return False
    lowered = text.lower()
    return any(term in lowered for term in _IMPORT_TERMS)


# ── 캐시 (성공 결과만) ───────────────────────────────────────────────────────
def _hash_targets() -> list:
    """캐시 키에 반영할 소스·데이터 파일. 이 파일들이 바뀌면 예전 캐시는 쓰이지 않는다."""
    data_dir = _SRC_DIR.parent / "data"
    return [
        _SRC_DIR / "main_pipeline.py",
        _SRC_DIR / "report_logic.py",
        _SRC_DIR / "hsk_inferencer.py",
        _SRC_DIR / "intent_router.py",
        _SRC_DIR / "tariff_catalog.py",
        _SRC_DIR / "rag_store.py",
        data_dir / "kc_exemption.txt",
        CSV_PATH,
    ]


def _normalized_hash(paths) -> str:
    """줄바꿈(CRLF/LF) 차이를 없앤 뒤 해시한다. Windows 와 Linux(클라우드)에서 같은 키가 나온다."""
    digest = hashlib.md5()
    for path in paths:
        try:
            digest.update(Path(path).read_bytes().replace(b"\r\n", b"\n"))
        except OSError:
            digest.update(b"missing")
        digest.update(b"\0")
    return digest.hexdigest()


def _cache_key(text: str, model_name: str, use_template: bool) -> str:
    normalized_query = " ".join(text.split())
    raw = f"{CACHE_VERSION}|{normalized_query}|{model_name}|{_normalized_hash(_hash_targets())}|tmpl={use_template}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()


def _cache_path(key: str) -> Path:
    return CACHE_DIR / f"{CACHE_PREFIX}{key}.json"


def _unlink(path: Path) -> None:
    try:
        path.unlink()
    except OSError:
        pass


def _restore_files(data: dict) -> None:
    """캐시 결과에 적힌 md/docx 가 없으면(서버 재시작 등) 다시 만든다."""
    md_path = Path(data.get("md_path", ""))
    docx_path = Path(data.get("docx_path", ""))
    report_md = data.get("report_markdown", "")
    try:
        if md_path.name and not md_path.exists():
            md_path.parent.mkdir(parents=True, exist_ok=True)
            md_path.write_text(report_md, encoding="utf-8")
        if docx_path.name and not docx_path.exists():
            from docx import Document

            doc = Document()
            doc.add_heading("R&D 패스트트랙 통관 및 세액 산출 보고서 (캐시에서 복구됨)", level=1)
            for line in report_md.split("\n"):
                doc.add_paragraph(line)
            docx_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(docx_path)
    except Exception as exc:  # noqa: BLE001 - 파일 복구 실패는 결과 반환을 막지 않는다
        logger.warning("캐시 결과 파일 복구 실패: %s", exc)


def _load_cache(path: Path):
    """성공 결과만 돌려준다. 읽을 수 없거나 성공이 아닌 파일은 지운다."""
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        logger.warning("캐시 읽기 실패로 파일을 삭제합니다: %s", exc)
        _unlink(path)
        return None
    if not isinstance(data, dict) or data.get("status") != "success":
        _unlink(path)
        return None
    _restore_files(data)
    data["cache_hit"] = True
    return data


def _save_cache(path: Path, data: dict) -> None:
    """성공한 결과(status == 'success')만 저장한다. 임시 파일에 쓴 뒤 교체해서 동시 접속에도 깨지지 않는다."""
    if data.get("status") != "success":
        return
    try:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        tmp = path.with_name(f"{path.name}.{uuid.uuid4().hex[:8]}.tmp")
        tmp.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        tmp.replace(path)
    except OSError as exc:
        logger.warning("캐시 저장 실패: %s", exc)


def clear_cache() -> int:
    """저장된 캐시를 모두 지우고 삭제한 개수를 돌려준다 (예전 버전 캐시 포함)."""
    removed = 0
    if CACHE_DIR.exists():
        for path in list(CACHE_DIR.glob("cache_*.json")) + list(CACHE_DIR.glob("*.tmp")):
            try:
                path.unlink()
                removed += 1
            except OSError:
                pass
    return removed


# ── 실행 ────────────────────────────────────────────────────────────────────
def run_pipeline(query: str, use_template: bool = True) -> dict:
    text = str(query or "").strip()
    if not text:
        return _fail("input", "질문을 입력해 주세요.")

    model_name = getattr(config, "MODEL_NAME", "gemini-3.8-flash")
    cache_file = _cache_path(_cache_key(text, model_name, use_template))
    cached = _load_cache(cache_file)
    if cached is not None:
        logger.info("캐시된 성공 응답을 반환합니다.")
        return cached

    qty = rl.extract_quantity(text)

    # ── 1. 의도 분류: 수입 관련 표현 + 수량이면 규칙으로, 그 밖은 LLM 라우터로 ──
    if _looks_like_import_request(text, qty):
        route_info = {
            "category": CUSTOMS,
            "confidence": 1.0,
            "reason": "규칙 기반 식별(수입 관련 표현 + 수량 확인됨)",
            "fallback_used": False,
        }
    else:
        try:
            router = _get_router()
        except RuntimeError as exc:  # API 키 미설정
            return _fail("router", str(exc), error_kind="config")

        try:
            route = router.route(text)
        except Exception as exc:  # noqa: BLE001
            logger.exception("의도 분류 실패")
            return _fail("router", f"의도 분류 중 오류가 발생했습니다 ({type(exc).__name__}).")

        route_info = {
            "category": route.category,
            "confidence": route.confidence,
            "reason": route.reason,
            "fallback_used": route.fallback_used,
        }
        if route.fallback_used and route.category == UNKNOWN:
            # LLM 이 실패했고 키워드도 맞지 않은 경우: '관련 없는 질문'이 아니라 장애 안내를 보여 준다
            message = route.reason.split(": ", 1)[-1]
            return _fail("router", message, error_kind=_kind_from_message(message), route=route_info)

    if route_info["category"] == LAB_SAFETY:
        return _fail("router", "연구실 안전 문의는 아직 지원하지 않습니다(Phase 2 예정).", status="unsupported", route=route_info)
    if route_info["category"] != CUSTOMS:
        return _fail("router", "수입 통관·인증과 관련된 질문이 아닙니다.", status="unsupported", route=route_info)

    # ── 2. 규칙 검사 (통관 질문으로 확정된 뒤, HSK 추론 앞): 거절될 입력에 LLM 호출을 쓰지 않는다 ──
    if qty is None:
        return _fail("report", QTY_MESSAGE, route=route_info)
    if rl.detect_purpose(text) == "commercial":
        return _fail("report", getattr(rl, "COMMERCIAL_NOTICE", "상업적 판매 목적의 수입은 지원 대상이 아닙니다."), route=route_info)
    if rl.extract_unit_price(text, qty).get("status") == "foreign":
        return _fail("report", FOREIGN_PRICE_MESSAGE, route=route_info)

    # ── 3. HSK 추론 ──
    infer = infer_hsk_code(text)
    if not isinstance(infer, dict) or infer.get("status") != "success":
        message = infer.get("message") if isinstance(infer, dict) else None
        extra = {}
        if isinstance(infer, dict) and infer.get("error_kind"):
            extra["error_kind"] = infer["error_kind"]
        return _fail("hsk", message or "HSK 추론에 실패했습니다.", route=route_info, **extra)

    data = infer.get("data")
    if not isinstance(data, dict):
        return _fail("hsk", "HSK 추론 결과 형식이 올바르지 않습니다.", route=route_info)

    hsk_code = str(data.get("hsk_code") or "").strip()
    inference_info = {
        "hsk_code": hsk_code,
        "item_name": data.get("item_name"),
        "confidence": data.get("confidence"),
        "reason": data.get("reason"),
    }
    if not hsk_code:
        reason = data.get("reason") or "근거 없음"
        return _fail("hsk", f"지원 대상 품목을 특정하지 못했습니다. (사유: {reason})", route=route_info, inference=inference_info)

    # ── 4. 보고서 생성 (템플릿 모드면 이 요청의 사유서 LLM 호출만 건너뜀) ──
    try:
        with _template_mode(use_template):
            report = generate_fast_track_report(text, hsk_hint=hsk_code)
    except Exception as exc:  # noqa: BLE001
        logger.exception("보고서 생성 실패")
        return _fail("report", f"보고서 생성 중 오류가 발생했습니다 ({type(exc).__name__}).", route=route_info, inference=inference_info)

    if not isinstance(report, dict) or report.get("status") != "success":
        message = report.get("message") if isinstance(report, dict) else None
        return _fail("report", message or "보고서를 생성하지 못했습니다.", route=route_info, inference=inference_info)

    result = dict(report)
    result.update({"stage": "done", "route": route_info, "inference": inference_info})
    _save_cache(cache_file, result)  # 성공 결과만 저장
    return result


if __name__ == "__main__":
    sample = "연구개발 목적으로 리튬이온 배터리 프로토타입 25대를 수입하려고 해."
    res = run_pipeline(" ".join(sys.argv[1:]) or sample)
    print(f"status={res['status']} stage={res.get('stage', '')}")
    if res["status"] == "success":
        print("MD   :", res.get("md_path"))
        print("DOCX :", res.get("docx_path"))
    else:
        print(res["message"])
