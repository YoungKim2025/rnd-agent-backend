"""R&D 패스트트랙 보고서 생성 (URL 링크 및 법령 근거 강화 버전)"""
import datetime
import re
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import config  # noqa: E402
from report_logic import (  # noqa: E402,F401
    COMMERCIAL_NOTICE,
    DEFAULT_UNIT_PRICE_KRW,
    UNREGISTERED_LIMIT_MESSAGE,
    build_tax_scenarios,
    calculate_estimated_tax,
    check_quantity,
    detect_purpose,
    extract_quantity,
    extract_unit_price,
    find_quantities,
    is_qc_purpose,
    mentions_origin_cert,
)
from tariff_catalog import lookup_tariff_strict  # noqa: E402

ROOT_DIR = config.ROOT_DIR
CSV_PATH = ROOT_DIR / "data" / "tariff_rates.csv"
CHROMA_DIR = config.CHROMA_DIR
OUTPUT_DIR = config.OUTPUT_DIR

DISCLAIMER = "본 결과는 참고용 초안입니다. 실제 신청 전 관세사 또는 소관 기관(국립전파연구원, KTC 등)에 확인하세요."
MAX_REASON_CHARS = 900
PLEDGE_SENTENCE = ("신청인은 본 물품을 상업적 판매·유통 목적으로 사용하지 않으며, 면제받은 물품의 유통기록과 "
                   "연구개발 사용 내역에 대한 사후 추적 관리를 이행할 것을 서약합니다.")
_ARTICLE_RE = re.compile(r"제\s*\d+\s*조(?:\s*의\s*\d+)?")

# 💡 신규: 법령별 공식 메타데이터 및 원문 URL 매핑
LAW_REFERENCES = {
    "전안법": [
        "「전기용품 및 생활용품 안전관리법」 본문 [시행 2025. 10. 1.] (제6조 안전인증의 면제 등)\n  🔗 https://www.law.go.kr/LSW/lsInfoP.do?lsiSeq=276591",
        "「전기용품 및 생활용품 안전관리법 시행령」 [시행 2026. 2. 1.] (제8조 안전인증의 면제)\n  🔗 https://law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1032807261"
    ],
    "전파법": [
        "「전파법」 본문 (제58조의2, 제58조의3 등)\n  🔗 https://www.law.go.kr/lsInfoP.do?lsId=001732",
        "「전파법 시행령」 [별표 6의2] 적합성평가 면제대상 기자재\n  🔗 https://www.law.go.kr/LSW/flDownload.do?bylClsCd=110201&flSeq=157120761"
    ]
}

def get_law_reference_links(law_type: str) -> list:
    """품목의 law_type에 맞춰 삽입할 URL 리스트를 반환합니다."""
    links = []
    if law_type in ("전파법", "복합"):
        links.extend(LAW_REFERENCES["전파법"])
    if law_type in ("전안법", "복합"):
        links.extend(LAW_REFERENCES["전안법"])
    return links

def lookup_tariff_by_hsk(hsk_code: str) -> dict:
    return lookup_tariff_strict(hsk_code)

def won(value: int) -> str:
    return f"₩{value:,}"

# ── RAG ─────────────────────────────────────────────────────────────────────
def retrieve_context(hsk_code: str, law_type_kr: str):
    from rag_store import format_context, retrieve_legal_docs, source_labels
    # 임베딩이 제거되었으므로 query 파라미터는 빈 문자열로 전달합니다.
    docs = retrieve_legal_docs("", hsk_code, law_type_kr)
    return format_context(docs), source_labels(docs)

# ── 사유서 본문 ─────────────────────────────────────────────────────────────
def _fallback_reason(item_name: str, quantity: int, qty_check: dict) -> str:
    text = (
        f"본 수입 물품({item_name}) {quantity:,}대는 당사의 연구개발(R&D) 및 프로토타입 성능 시험을 위해 반입하는 시제품입니다. "
        + PLEDGE_SENTENCE
    )
    if qty_check["state"] == "within":
        text += " " + qty_check["message"]
    return text

def _has_pledge(text: str) -> bool:
    return ("서약" in text or "확약" in text) and "추적" in text

def _ensure_pledge(text: str) -> str:
    return text if _has_pledge(text) else f"{text} {PLEDGE_SENTENCE}"

def _append_limit_note(text: str, qty_check: dict) -> str:
    if qty_check["state"] == "unregistered" and UNREGISTERED_LIMIT_MESSAGE not in text:
        return f"{text} ※ {UNREGISTERED_LIMIT_MESSAGE}"
    return text

def _reason_is_safe(text: str, law_basis: str, qty_check: dict, law_type: str) -> bool:
    if not text or len(text) > MAX_REASON_CHARS or re.search(r"^\s*([-*•]|\d+\.)\s", text, re.MULTILINE):
        return False
    if qty_check["state"] != "within" and "충족" in re.sub(r"충족\s*여부", "", text):
        return False  
    
    if law_type in ("전안법", "복합"):
        if "제9조" in text or "제11조" in text:
            return False

    allowed = {re.sub(r"\s", "", a) for a in _ARTICLE_RE.findall(law_basis)}
    cited = {re.sub(r"\s", "", a) for a in _ARTICLE_RE.findall(text)}
    return cited <= allowed

def generate_rich_reason(item_name: str, quantity: int, law_basis: str, rag_context: str, qty_check: dict, law_type: str) -> str:
    prompt = f"""
당신은 관세법인의 전문 기안자입니다. 아래 [검증된 사실]만 근거로 소관 기관에 제출할 '요건 면제 사유서' 본문을
3~4문장의 비즈니스 톤 한 문단으로 작성해 주세요.

[검증된 사실]
- 수입 품목: {item_name}
- 수량: {quantity}대
- 수입 목적: 연구개발(R&D) 및 프로토타입 시험 (상업적 판매 목적 아님)
- 수량 한도 검증 결과: {qty_check['message']}
- 표기할 법령: {law_basis}

[참고 규정 발췌 - 문맥 이해용]
{rag_context or '(검색된 규정 없음)'}

[작성 규칙]
1. [검증된 사실]에 없는 내용을 사실처럼 단정하지 마세요. 특히 '수량 기준 충족'은 위 검증 결과가 '한도 이내'일 때만 쓰세요.
2. 법령 표기는 위 [표기할 법령]의 문구만 쓰고, 조문 번호를 새로 만들지 마세요. (전안법의 경우 제9조, 제11조는 면제 근거가 아니므로 절대 인용하지 마세요)
3. 신청인이 서약하는 사항(상업적 판매·유통 금지, 유통기록과 연구개발 사용 내역의 사후 추적 관리)을 포함하세요.
4. 마크다운·불릿 없이 순수한 한 문단만 출력하세요.
"""
    try:
        text = config.generate_text(prompt)
    except Exception as exc:  
        print(f"[경고] 사유서 LLM 호출 실패, 표준 문구 사용: {type(exc).__name__}: {str(exc)[:200]}")
        return _append_limit_note(_ensure_pledge(_fallback_reason(item_name, quantity, qty_check)), qty_check)
    draft = text if _reason_is_safe(text, law_basis, qty_check, law_type) else _fallback_reason(item_name, quantity, qty_check)
    return _append_limit_note(_ensure_pledge(draft), qty_check)

# ── 문서 생성 ───────────────────────────────────────────────────────────────
def _scenario_rows(tariff_info: dict, tax: dict) -> list:
    applied = tax["applied_key"]
    return [
        ("기본세율", tariff_info["base_rate"], tax["base"], applied == "base", ""),
        ("FTA 협정세율", tariff_info["fta_rate"], tax["fta"], applied == "fta", "원산지증명서·원산지 결정 기준 충족 시"),
    ]

def build_report_markdown(*, query, tariff_info, quantity, purpose, tax, price_note, qty_check,
                          reason, sources, warnings) -> str:
    warning_block = ""
    if warnings:
        warning_block = "> ⚠️ **확인이 필요한 사항**\n" + "\n".join(f"> - {w}" for w in warnings) + "\n\n"

    purpose_label = {"research": "연구개발(R&D) — 입력 기준", "unknown": "확인되지 않음"}.get(purpose, purpose)
    rows = "\n".join(
        f"| {name}{' ◀ 적용' if applied else ''} | {rate} | {won(t['estimated_duty'])} | {won(t['estimated_vat'])} | **{won(t['total_tax'])}** |"
        for name, rate, t, applied, _note in _scenario_rows(tariff_info, tax)
    )
    if purpose == "research":
        exemption_md = f"**[문서 제목] {tariff_info['law_basis']} 신청 사유서**\n\n> {reason}\n"
    else:
        exemption_md = ("> 수입 목적(연구개발 등)이 입력에서 확인되지 않아 면제 사유서 본문을 작성하지 않았습니다. "
                        "수입 목적을 포함해 다시 입력해 주세요.\n")
                        
    source_md = "\n".join(f"- {s}" for s in sources) if sources else "- (참조된 사내 규정 문서 없음)"
    law_links = get_law_reference_links(tariff_info.get("law_type", ""))
    if law_links:
        source_md += "\n\n**[국가법령정보센터 공식 원문 및 URL]**\n"
        source_md += "\n".join(f"- {link}" for link in law_links)

    return f"""## R&D 패스트트랙 통관 및 세액 산출 보고서

> **[사용자 원본 요청]** "{query}"

{warning_block}---
### 1. 🔍 품목 분류 및 수량 분석 결과
- **식별된 품목**: {tariff_info['item_name']}
- **추론 HSK 코드**: **{tariff_info['hsk_code']}**
- **요청 수량**: {quantity:,}대
- **수입 목적**: {purpose_label}
- **수량 한도 검증**: {qty_check['message']}
> ℹ️ [데이터 소스] {tariff_info.get('source', 'local')}

---
### 2. 📊 예상 관세율표 및 세액 산출
- **적용 과세가격**: 단가 {won(tax['unit_krw'])} × {quantity:,}대 = **총 {won(tax['total_krw'])}** ({price_note})

| 구분 | 세율 | 관세 | 부가세(10%) | 총 예상 세액 |
| :--- | :---: | :---: | :---: | :---: |
{rows}

> 부가세 = (과세가격 + 관세) × 10%. FTA 세율은 원산지증명서(C/O)와 협정별 원산지 결정 기준 충족이 전제이며,
> 사용자가 C/O 보유를 밝히지 않으면 **기본세율**을 적용세율로 봅니다.

---
### 3. 📄 인증 면제 사유서 초안
{exemption_md}
---
### 4. 📎 필수 구비 서류 (관세청 UNI-PASS 제출용)
1. 면제 확인(신청)서
2. 면제사실 증명 서류 (연구개발계획서, 기술개발계획서, 시제품 활용 목적 사유서 등)
3. 물품 확인 서류 (화물송장 Commercial Invoice, 수입계약서 등)

---
### 5. 📚 참조한 규정 및 법령 원문
{source_md}

> {DISCLAIMER}
"""

def save_to_docx(*, query, tariff_info, quantity, purpose, tax, price_note, qty_check,
                 reason, sources, warnings, output_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("R&D 패스트트랙 통관 및 세액 산출 보고서", level=1)

    p_req = doc.add_paragraph()
    p_req.add_run("원본 요청: ").bold = True
    p_req.add_run(f'"{query}"')

    if warnings:
        doc.add_heading("확인이 필요한 사항", level=2)
        for w in warnings:
            doc.add_paragraph(w, style="List Bullet")

    doc.add_heading("1. 품목 분류 및 수량 분석", level=2)
    doc.add_paragraph(f"식별 품목: {tariff_info['item_name']}", style="List Bullet")
    doc.add_paragraph(f"HSK 코드: {tariff_info['hsk_code']}", style="List Bullet")
    doc.add_paragraph(f"요청 수량: {quantity:,}대", style="List Bullet")
    doc.add_paragraph(f"수량 한도 검증: {qty_check['message']}", style="List Bullet")

    doc.add_heading("2. 예상 관세 및 세액 산출", level=2)
    doc.add_paragraph(f"과세가격: 단가 {won(tax['unit_krw'])} × {quantity:,}대 = {won(tax['total_krw'])} ({price_note})")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("구분", "세율", "관세", "부가세(10%)", "총 예상 세액")):
        cell.text = text
    for name, rate, t, applied, _ in _scenario_rows(tariff_info, tax):
        cells = table.add_row().cells
        cells[0].text = name + (" (적용)" if applied else "")
        cells[1].text = rate
        cells[2].text = won(t["estimated_duty"])
        cells[3].text = won(t["estimated_vat"])
        cells[4].text = won(t["total_tax"])
    doc.add_paragraph("FTA 세율은 원산지증명서(C/O)와 원산지 결정 기준 충족이 전제이며, C/O 보유를 밝히지 않은 경우 기본세율을 적용세율로 봅니다.")

    doc.add_heading("3. 인증 면제 사유서 본문", level=2)
    if purpose == "research":
        doc.add_paragraph(tariff_info["law_basis"], style="Intense Quote")
        doc.add_paragraph(reason)
    else:
        doc.add_paragraph("수입 목적이 입력에서 확인되지 않아 면제 사유서 본문을 작성하지 않았습니다.")

    doc.add_heading("4. 필요 증빙 서류 (UNI-PASS 첨부용)", level=2)
    for item in ("면제 확인(신청)서", "연구개발계획서 및 시제품 활용 사유서", "화물송장(Commercial Invoice)"):
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("5. 참조한 규정 및 법령 원문", level=2)
    for source in sources or ["(참조된 사내 규정 문서 없음)"]:
        doc.add_paragraph(source, style="List Bullet")

    law_links = get_law_reference_links(tariff_info.get("law_type", ""))
    if law_links:
        p_law = doc.add_paragraph("국가법령정보센터 공식 원문 및 URL")
        p_law.runs[0].bold = True
        for link in law_links:
            parts = link.split("\n  🔗 ")
            p = doc.add_paragraph(parts[0], style="List Bullet")
            if len(parts) > 1:
                p.add_run(f"\nURL: {parts[1]}")

    note = doc.add_paragraph(DISCLAIMER)
    note.runs[0].italic = True
    doc.save(output_path)


# ── 진입점 ──────────────────────────────────────────────────────────────────
def generate_fast_track_report(query: str, hsk_hint: str) -> dict:
    warnings = []

    quantities = find_quantities(query)
    if not quantities:
        return {"status": "error", "message": "수량을 확인하지 못했습니다. 예: '배터리 25개'처럼 수량을 포함해 다시 입력해 주세요."}
    quantity = quantities[0]
    if len(quantities) > 1:
        warnings.append(f"수량 후보가 여러 개 감지되어({', '.join(f'{q:,}' for q in quantities)}) 첫 번째 값 {quantity:,}대를 사용했습니다.")

    purpose = detect_purpose(query)
    if purpose == "commercial":
        return {"status": "error", "message": COMMERCIAL_NOTICE}

    tariff_info = lookup_tariff_by_hsk(hsk_hint)
    if tariff_info.get("status") != "success":
        return {"status": "error", "message": tariff_info.get("message", "관세율 정보를 찾지 못했습니다.")}
    tariff_info = dict(tariff_info)
    tariff_info["law_basis"] = tariff_info.get("law_basis") or tariff_info["law_title"]

    price = extract_unit_price(query, quantity)
    if price["status"] == "foreign":
        return {"status": "error", "message": "외화 단가가 입력되었습니다. 적용 환율을 알 수 없어 계산하지 않았습니다. 원화로 환산한 단가(예: '개당 130만원')로 다시 입력해 주세요."}
    if price["status"] == "found":
        unit_price = price["unit_krw"]
        price_note = {"unit": "사용자 입력 단가", "total": "입력한 총액을 수량으로 나눈 단가", "unspecified": "입력 금액을 단가로 해석"}[price["basis"]]
        if price["basis"] == "unspecified":
            warnings.append("입력된 금액이 단가인지 총액인지 명확하지 않아 단가로 해석했습니다. ('개당 ○○원'처럼 적으면 정확합니다)")
    else:
        unit_price = DEFAULT_UNIT_PRICE_KRW
        price_note = "단가 미입력 — 가정값"
        warnings.append(f"단가가 입력되지 않아 가정값 {won(DEFAULT_UNIT_PRICE_KRW)}으로 계산했습니다. 세액은 실제 단가(CIF 기준 과세가격)로 다시 산출해야 합니다.")

    qty_check = check_quantity(quantity, tariff_info.get("limit_quantity", 0), tariff_info.get("law_type", ""), is_qc_purpose(query))
    if qty_check["state"] == "exceeded":
        return {"status": "error", "message": qty_check["message"]}
    if qty_check["state"] == "unregistered":
        warnings.append(qty_check["message"])

    fta_applied = mentions_origin_cert(query)
    if fta_applied:
        warnings.append("원산지증명서 보유로 입력되어 FTA 협정세율을 적용했습니다. 협정별 원산지 결정 기준 충족 여부는 별도 확인이 필요합니다.")
    tax = build_tax_scenarios(tariff_info["base_rate_val"], tariff_info["fta_rate_val"], quantity, unit_price, fta_applied)

    reason, sources, rag_context = "", [], ""
    if purpose == "research":
        try:
            # 💡 수정됨: 더 이상 임베딩 검색에 사용하지 않는 query 전달 생략
            rag_context, sources = retrieve_context(tariff_info["hsk_code"], tariff_info.get("law_type", ""))
            if not rag_context:
                warnings.append("관련 규정 조항을 찾지 못해 규정 발췌 없이 사유서를 작성했습니다.")
        except Exception as exc: 
            print(f"[경고] 규정 검색 실패: {type(exc).__name__}: {exc}")
            warnings.append("규정 검색에 실패해 규정 발췌 없이 사유서를 작성했습니다. (벡터 스토어를 확인하세요)")
        
        reason = generate_rich_reason(tariff_info["item_name"], quantity, tariff_info["law_basis"], rag_context, qty_check, tariff_info.get("law_type", ""))
    else:
        warnings.append("수입 목적(연구개발 등)이 입력에서 확인되지 않아 면제 사유서 본문은 작성하지 않았습니다.")

    fields = dict(query=query, tariff_info=tariff_info, quantity=quantity, purpose=purpose, tax=tax,
                  price_note=price_note, qty_check=qty_check, reason=reason, sources=sources, warnings=warnings)
    report_md = build_report_markdown(**fields)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = f"{datetime.datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}" 
    md_path = OUTPUT_DIR / f"tariff_report_{stamp}.md"
    docx_path = OUTPUT_DIR / f"tariff_report_{stamp}.docx"
    md_path.write_text(report_md, encoding="utf-8")
    save_to_docx(output_path=docx_path, **fields)

    return {
        "status": "success",
        "report_markdown": report_md,
        "md_path": str(md_path),
        "docx_path": str(docx_path),
        "warnings": warnings,
        "quantity": quantity,
        "unit_price_krw": unit_price,
        "tax": tax,
    }


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    from pipeline_runner import run_pipeline

    result = run_pipeline(" ".join(sys.argv[1:]) or "연구개발 목적으로 리튬이온 배터리 프로토타입 25개를 수입하려고 해. 개당 5천원이야.")
    print(f"status={result['status']} stage={result.get('stage')}")
    if result["status"] == "success":
        for w in result.get("warnings", []):
            print("[확인 필요]", w)
        print("MD   :", result["md_path"])
        print("DOCX :", result["docx_path"])
    else:
        print(result["message"])